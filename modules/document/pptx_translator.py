"""PPTX presentation translator"""
from pathlib import Path
from typing import Optional


class PPTXTranslator:
    """Translate PowerPoint presentations"""
    
    def __init__(self, translator):
        self.translator = translator
    
    def translate_file(self, input_path: str, output_path: str,
                      source_lang: str, target_lang: str,
                      glossary: Optional[dict] = None) -> bool:
        """
        Translate PPTX file
        
        Args:
            input_path: Input PPTX file path
            output_path: Output PPTX file path
            source_lang: Source language code
            target_lang: Target language code
            glossary: Optional glossary dictionary
            
        Returns:
            True if successful
        """
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("python-pptx not installed. Install with: pip install python-pptx")
        
        try:
            # Load presentation
            prs = Presentation(input_path)
            
            # Translate each slide
            for slide in prs.slides:
                # Translate shapes with text
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        original_text = shape.text
                        
                        # Apply glossary
                        if glossary:
                            for source_term, target_term in glossary.items():
                                original_text = original_text.replace(source_term, target_term)
                        
                        # Translate
                        translated_text = self.translator.translate(original_text, source_lang, target_lang)
                        
                        # Update text frame
                        if hasattr(shape, "text_frame"):
                            shape.text_frame.clear()
                            p = shape.text_frame.paragraphs[0]
                            p.text = translated_text
                    
                    # Translate table cells
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            for cell in row.cells:
                                if cell.text.strip():
                                    original_text = cell.text
                                    
                                    if glossary:
                                        for source_term, target_term in glossary.items():
                                            original_text = original_text.replace(source_term, target_term)
                                    
                                    translated_text = self.translator.translate(original_text, source_lang, target_lang)
                                    cell.text = translated_text
            
            # Save translated presentation
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            prs.save(output_path)
            return True
            
        except Exception as e:
            print(f"PPTX translation error: {e}")
            return False
