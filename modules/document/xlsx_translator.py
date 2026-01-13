"""XLSX spreadsheet translator"""
from pathlib import Path
from typing import Optional


class XLSXTranslator:
    """Translate Excel spreadsheets"""
    
    def __init__(self, translator):
        self.translator = translator
    
    def translate_file(self, input_path: str, output_path: str,
                      source_lang: str, target_lang: str,
                      glossary: Optional[dict] = None) -> bool:
        """
        Translate XLSX file
        
        Args:
            input_path: Input XLSX file path
            output_path: Output XLSX file path
            source_lang: Source language code
            target_lang: Target language code
            glossary: Optional glossary dictionary
            
        Returns:
            True if successful
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("openpyxl not installed. Install with: pip install openpyxl")
        
        try:
            # Load workbook
            wb = load_workbook(input_path)
            
            # Translate each sheet
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                
                # Translate cells
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and cell.value.strip():
                            original_text = cell.value
                            
                            # Apply glossary
                            if glossary:
                                for source_term, target_term in glossary.items():
                                    original_text = original_text.replace(source_term, target_term)
                            
                            # Translate
                            try:
                                translated_text = self.translator.translate(original_text, source_lang, target_lang)
                                cell.value = translated_text
                            except Exception as e:
                                print(f"Cell translation error: {e}")
                                # Keep original value if translation fails
                                pass
            
            # Save translated workbook
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            return True
            
        except Exception as e:
            print(f"XLSX translation error: {e}")
            return False
