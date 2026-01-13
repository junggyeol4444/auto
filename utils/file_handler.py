"""
Utility module for file operations
"""
import os
import json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import PyPDF2


class FileHandler:
    """Handle file I/O operations for various formats"""
    
    @staticmethod
    def read_text(file_path):
        """Read text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            raise Exception(f"텍스트 파일 읽기 오류: {str(e)}")
    
    @staticmethod
    def read_docx(file_path):
        """Read text from DOCX file"""
        try:
            doc = Document(file_path)
            return '\n'.join([para.text for para in doc.paragraphs])
        except Exception as e:
            raise Exception(f"DOCX 파일 읽기 오류: {str(e)}")
    
    @staticmethod
    def write_text(file_path, content):
        """Write text to TXT file"""
        try:
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            return True
        except Exception as e:
            raise Exception(f"텍스트 파일 저장 오류: {str(e)}")
    
    @staticmethod
    def write_docx(file_path, content, title=None, chapters=None):
        """Write text to DOCX file with formatting"""
        try:
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            doc = Document()
            
            # Add title if provided
            if title:
                heading = doc.add_heading(title, 0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add chapters if structured
            if chapters:
                for chapter_num, chapter_content in enumerate(chapters, 1):
                    doc.add_heading(f'제{chapter_num}화', 1)
                    for paragraph in chapter_content.split('\n\n'):
                        if paragraph.strip():
                            p = doc.add_paragraph(paragraph.strip())
                            p.paragraph_format.line_spacing = 1.5
                            p.paragraph_format.space_after = Pt(12)
            else:
                # Add content as is
                for paragraph in content.split('\n\n'):
                    if paragraph.strip():
                        p = doc.add_paragraph(paragraph.strip())
                        p.paragraph_format.line_spacing = 1.5
                        p.paragraph_format.space_after = Pt(12)
            
            doc.save(file_path)
            return True
        except Exception as e:
            raise Exception(f"DOCX 파일 저장 오류: {str(e)}")
    
    @staticmethod
    def write_manual_docx(file_path, title, sections):
        """Write manual to DOCX with proper formatting"""
        try:
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            doc = Document()
            
            # Title
            title_heading = doc.add_heading(title, 0)
            title_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Table of Contents placeholder
            doc.add_heading('목차', 1)
            doc.add_paragraph('[자동 생성된 목차]')
            doc.add_page_break()
            
            # Add sections
            for section in sections:
                doc.add_heading(section['title'], 1)
                for subsection in section.get('subsections', []):
                    if isinstance(subsection, dict):
                        doc.add_heading(subsection.get('title', ''), 2)
                        content = subsection.get('content', '')
                        for para in content.split('\n'):
                            if para.strip():
                                doc.add_paragraph(para.strip())
                    else:
                        doc.add_paragraph(subsection, style='List Bullet')
                
                if 'content' in section:
                    for para in section['content'].split('\n'):
                        if para.strip():
                            doc.add_paragraph(para.strip())
            
            doc.save(file_path)
            return True
        except Exception as e:
            raise Exception(f"매뉴얼 DOCX 저장 오류: {str(e)}")
    
    @staticmethod
    def load_json(file_path):
        """Load JSON file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            raise Exception(f"JSON 파일 로드 오류: {str(e)}")
    
    @staticmethod
    def save_json(file_path, data):
        """Save to JSON file"""
        try:
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            raise Exception(f"JSON 파일 저장 오류: {str(e)}")
